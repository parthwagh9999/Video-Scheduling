"""
Unit tests for Excel synchronization and validation.
"""
import pytest
from pathlib import Path
import openpyxl
from config.settings import EXCEL_COLUMNS, EXCEL_SHEET_NAME
from app.excel.sync import (
    validate_workbook_structure, read_all_records,
    find_missing_videos, create_backup, preview_new_records
)


@pytest.fixture
def mock_workbook(tmp_path):
    """Create a temporary excel file matching the real workbook schema."""
    wb_path = tmp_path / "test_seo.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = EXCEL_SHEET_NAME

    # Headers
    for col_idx, col_name in enumerate(EXCEL_COLUMNS, start=1):
        ws.cell(row=1, column=col_idx, value=col_name)

    # 2 Sample Rows
    ws.cell(row=2, column=1, value="sample_01.mp4")
    ws.cell(row=2, column=2, value="Title 1 #shorts")
    ws.cell(row=2, column=3, value="Description 1")
    ws.cell(row=2, column=4, value=2)
    ws.cell(row=2, column=5, value="Pending")

    ws.cell(row=3, column=1, value="sample_02.mp4")
    ws.cell(row=3, column=2, value="Title 2 #shorts")
    ws.cell(row=3, column=3, value="Description 2")
    ws.cell(row=3, column=4, value=2)
    ws.cell(row=3, column=5, value="Scheduled")

    wb.save(str(wb_path))
    wb.close()
    return wb_path


def test_validate_workbook_structure(mock_workbook):
    wb = openpyxl.load_workbook(str(mock_workbook))
    res = validate_workbook_structure(wb)
    assert res['valid'] is True
    assert len(res['issues']) == 0
    wb.close()


def test_read_all_records(mock_workbook):
    records = read_all_records(mock_workbook)
    assert len(records) == 2
    assert records[0]['filename'] == "sample_01.mp4"
    assert records[0]['seo_title'] == "Title 1 #shorts"
    assert records[1]['status'] == "Scheduled"


def test_find_missing_videos(mock_workbook, tmp_path):
    video_dir = tmp_path / "videos"
    video_dir.mkdir()
    # Create sample_01.mp4 and sample_03.mp4
    (video_dir / "sample_01.mp4").write_text("fake video")
    (video_dir / "sample_03.mp4").write_text("fake video")

    res = find_missing_videos(mock_workbook, start_dir=video_dir)
    assert "sample_01.mp4" in res['matched']
    assert "sample_03.mp4" in res['missing_from_excel']
    assert "sample_02.mp4" in res['missing_from_disk']
